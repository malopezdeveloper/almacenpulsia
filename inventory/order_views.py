from django.contrib import messages
from django.contrib.auth.decorators import login_required, user_passes_test
from django.contrib.auth import get_user_model
from django.db import transaction
from django.http import HttpResponseForbidden
from django.shortcuts import get_object_or_404, redirect, render
from django.utils import timezone
from django.utils.text import slugify
from .models import InventoryRecord
from .order_models import BusinessRole,BusinessRoleAssignment,Customer,Supplier,CustomerOrder,OrderUnit,Component,Repair,ComponentReservation,RMA,ProcurementAlert
from .permissions import PERMISSION_CHOICES,user_has_permission,user_is_purchasing

def _allowed(user,perm): return user.is_staff or user_has_permission(user,perm)
def _deny(): return HttpResponseForbidden('No tienes permiso para realizar esta operación.')

@login_required
def orders_center(request):
 if not _allowed(request.user,'orders.view'): return _deny()
 ctx={'customers':Customer.objects.filter(active=True).order_by('name'),'suppliers':Supplier.objects.filter(active=True).order_by('name'),'orders':CustomerOrder.objects.select_related('customer').order_by('-id')[:100],'units':OrderUnit.objects.select_related('order','order__customer').order_by('-id')[:100],'components':Component.objects.select_related('supplier','inventory_record').order_by('-id')[:150],'repairs':Repair.objects.select_related('unit','created_by').order_by('-id')[:100],'reservations':ComponentReservation.objects.select_related('repair','component','technician').order_by('-id')[:100],'rmas':RMA.objects.select_related('component','supplier').order_by('-id')[:100],'alerts':ProcurementAlert.objects.filter(status='open').select_related('repair','repair__unit').order_by('-created_at') if user_is_purchasing(request.user) else ProcurementAlert.objects.none(),'inventory_records':InventoryRecord.objects.filter(status='available').select_related('table').order_by('table__name','internal_id')[:1000]}
 return render(request,'inventory/orders_center.html',ctx)

@login_required
def customer_create(request):
 if not _allowed(request.user,'customers.manage'): return _deny()
 if request.method=='POST': Customer.objects.create(name=request.POST['name'].strip(),phone=request.POST.get('phone','').strip(),email=request.POST.get('email','').strip(),address=request.POST.get('address','').strip(),delivery_point=request.POST.get('delivery_point','').strip(),contact=request.POST.get('contact','').strip(),observations=request.POST.get('observations','').strip()); messages.success(request,'Cliente creado.')
 return redirect('orders_center')
@login_required
def supplier_create(request):
 if not _allowed(request.user,'suppliers.manage'): return _deny()
 if request.method=='POST': Supplier.objects.create(name=request.POST['name'].strip(),phone=request.POST.get('phone','').strip(),email=request.POST.get('email','').strip(),address=request.POST.get('address','').strip(),delivery_point=request.POST.get('delivery_point','').strip(),contact=request.POST.get('contact','').strip(),observations=request.POST.get('observations','').strip()); messages.success(request,'Proveedor creado.')
 return redirect('orders_center')
@login_required
def order_create(request):
 if not _allowed(request.user,'orders.manage'): return _deny()
 if request.method=='POST': CustomerOrder.objects.create(name=request.POST['name'].strip(),customer_id=request.POST['customer'],brand=request.POST.get('brand','').strip(),model=request.POST.get('model','').strip(),lot=request.POST.get('lot','').strip(),processor=request.POST.get('processor','').strip(),ram=request.POST.get('ram','').strip(),disk=request.POST.get('disk','').strip(),created_by=request.user); messages.success(request,'Pedido creado.')
 return redirect('orders_center')
@login_required
def unit_create(request):
 if not _allowed(request.user,'orders.manage'): return _deny()
 if request.method=='POST':
  order=get_object_or_404(CustomerOrder,pk=request.POST['order']); OrderUnit.objects.create(order=order,serial_number=request.POST['serial_number'].strip(),aiken_lot=request.POST.get('aiken_lot','').strip() or order.lot,brand=request.POST.get('brand','').strip() or order.brand,model=request.POST.get('model','').strip() or order.model,processor=request.POST.get('processor','').strip() or order.processor,ram=request.POST.get('ram','').strip() or order.ram,disk=request.POST.get('disk','').strip() or order.disk); messages.success(request,'Unidad vinculada al pedido.')
 return redirect('orders_center')
@login_required
def component_create(request):
 if not _allowed(request.user,'components.manage'): return _deny()
 if request.method=='POST': Component.objects.create(component_type=request.POST['component_type'].strip(),supplier_id=request.POST.get('supplier') or None,reference=request.POST.get('reference','').strip(),inventory_record_id=request.POST.get('inventory_record') or None,observations=request.POST.get('observations','').strip()); messages.success(request,'Componente registrado.')
 return redirect('orders_center')
@login_required
def repair_create(request):
 if not _allowed(request.user,'repairs.manage'): return _deny()
 if request.method!='POST': return redirect('orders_center')
 with transaction.atomic():
  unit=get_object_or_404(OrderUnit,pk=request.POST['unit']); repair_type=request.POST['repair_type'].strip(); repair=Repair.objects.create(unit=unit,repair_type=repair_type,observations=request.POST.get('observations','').strip(),created_by=request.user); component=Component.objects.select_for_update().filter(component_type__iexact=repair_type,status='active').order_by('id').first()
  if component:
   ComponentReservation.objects.create(repair=repair,component=component,technician=request.user,unit_serial_number=unit.serial_number); component.status='reserved'; component.save(update_fields=['status'])
   if component.inventory_record_id:
    rec=component.inventory_record; rec.status='reserved'; rec.current_sn=unit.serial_number; rec.current_technician=request.user.get_username(); rec.save(update_fields=['status','current_sn','current_technician','updated_at'])
   messages.success(request,f'Reparación creada y componente {component} reservado automáticamente.')
  else: ProcurementAlert.objects.create(repair=repair,message=f'Falta componente {repair_type} para la unidad {unit.serial_number}.'); messages.error(request,f'ALERTA DE COMPRAS: no hay ningún componente {repair_type} disponible para {unit.serial_number}.')
 return redirect('orders_center')
@login_required
def reservation_cancel(request,pk):
 reservation=get_object_or_404(ComponentReservation,pk=pk)
 if request.method!='POST' or not (_allowed(request.user,'components.reserve') or reservation.technician_id==request.user.id): return _deny()
 if reservation.status=='active': reservation.cancel(); messages.success(request,'Reserva cancelada. El componente vuelve a estar activo/disponible.')
 return redirect('orders_center')
@login_required
def component_low(request,pk):
 if not _allowed(request.user,'components.manage'): return _deny()
 c=get_object_or_404(Component,pk=pk)
 if request.method=='POST':
  c.status='low'; c.save(update_fields=['status'])
  if c.inventory_record_id: r=c.inventory_record; r.status='scrapped'; r.save(update_fields=['status','updated_at'])
  messages.success(request,'Componente dado de baja. Ya puede tramitarse un RMA.')
 return redirect('orders_center')
@login_required
def rma_create(request):
 if not _allowed(request.user,'rma.manage'): return _deny()
 if request.method=='POST':
  c=get_object_or_404(Component,pk=request.POST['component'],status='low')
  if not c.supplier_id: messages.error(request,'El componente no tiene proveedor asociado.'); return redirect('orders_center')
  r=RMA(component=c,supplier=c.supplier,reason=request.POST.get('reason','').strip(),observations=request.POST.get('observations','').strip(),created_by=request.user); r.full_clean(); r.save(); messages.success(request,'RMA creado contra el proveedor del componente.')
 return redirect('orders_center')
@login_required
def procurement_resolve(request,pk):
 if not user_has_permission(request.user,'procurement.resolve') and not request.user.is_superuser: return _deny()
 a=get_object_or_404(ProcurementAlert,pk=pk)
 if request.method=='POST': a.status='resolved'; a.resolved_at=timezone.now(); a.save(update_fields=['status','resolved_at'])
 return redirect('orders_center')
@login_required
@user_passes_test(lambda u:u.is_superuser)
def roles_manager(request):
 User=get_user_model()
 if request.method=='POST':
  action=request.POST.get('action')
  if action=='create': name=request.POST['name'].strip(); BusinessRole.objects.create(name=name,code=slugify(name),permissions=[p for p,_ in PERMISSION_CHOICES if p in request.POST.getlist('permissions')]); messages.success(request,'Rol creado.')
  elif action=='update': role=get_object_or_404(BusinessRole,pk=request.POST['role']); role.permissions=[p for p,_ in PERMISSION_CHOICES if p in request.POST.getlist('permissions')]; role.active=request.POST.get('active')=='on'; role.save(); messages.success(request,'Permisos actualizados.')
  elif action=='assign':
   user=get_object_or_404(User,pk=request.POST['user']); BusinessRoleAssignment.objects.filter(user=user).delete()
   for rid in request.POST.getlist('roles'): BusinessRoleAssignment.objects.create(user=user,role_id=rid)
   messages.success(request,'Roles del usuario actualizados.')
  return redirect('roles_manager')
 return render(request,'inventory/roles_manager.html',{'roles':BusinessRole.objects.order_by('name'),'users':User.objects.filter(is_active=True).order_by('username'),'permission_choices':PERMISSION_CHOICES})

@login_required
@user_passes_test(lambda u:u.is_superuser)
def import_router(request):
 from . import views as legacy_views
 if request.method=='POST':
  destination=request.POST.get('destination','inventory')
  if destination=='inventory': return legacy_views.import_view(request)
  upload=request.FILES.get('file')
  if not upload: messages.error(request,'Seleccione un Excel.'); return redirect('import_excel')
  from openpyxl import load_workbook
  wb=load_workbook(upload,read_only=True,data_only=True); ws=wb.active; rows=list(ws.iter_rows(values_only=True)); headers=[str(x or '').strip().lower() for x in (rows[0] if rows else [])]
  def val(row,*names):
   for n in names:
    if n in headers: v=row[headers.index(n)]; return '' if v is None else str(v).strip()
   return ''
  count=0
  if destination=='customers':
   for row in rows[1:]:
    name=val(row,'nombre','cliente','name')
    if name: Customer.objects.update_or_create(name=name,defaults={'phone':val(row,'telefono','teléfono','phone'),'email':val(row,'email','correo'),'address':val(row,'direccion','dirección','address'),'delivery_point':val(row,'punto de entrega','punto_entrega'),'contact':val(row,'contacto'),'observations':val(row,'observaciones')}); count+=1
  elif destination=='suppliers':
   for row in rows[1:]:
    name=val(row,'nombre','proveedor','name')
    if name: Supplier.objects.update_or_create(name=name,defaults={'phone':val(row,'telefono','teléfono','phone'),'email':val(row,'email','correo'),'address':val(row,'direccion','dirección','address'),'delivery_point':val(row,'punto de entrega','punto_entrega'),'contact':val(row,'contacto'),'observations':val(row,'observaciones')}); count+=1
  elif destination=='units':
   order=get_object_or_404(CustomerOrder,pk=request.POST.get('order'))
   for row in rows[1:]:
    sn=val(row,'sn','serial','serial number','numero de serie','número de serie')
    if sn: OrderUnit.objects.update_or_create(serial_number=sn,defaults={'order':order,'aiken_lot':val(row,'lote','lot') or order.lot,'brand':val(row,'marca','brand') or order.brand,'model':val(row,'modelo','model') or order.model,'processor':val(row,'procesador','processor','cpu') or order.processor,'ram':val(row,'ram','memoria') or order.ram,'disk':val(row,'disco','disk','almacenamiento') or order.disk}); count+=1
  messages.success(request,f'Importación completada: {count} registros.'); return redirect('import_excel')
 return render(request,'inventory/import_router.html',{'orders':CustomerOrder.objects.order_by('-id')})

@login_required
@user_passes_test(lambda u:u.is_superuser)
def export_router(request):
 from . import views as legacy_views
 source=request.GET.get('source')
 if not source: return render(request,'inventory/export_router.html')
 if source=='inventory': return legacy_views.export_view(request)
 from openpyxl import Workbook
 from django.http import HttpResponse
 wb=Workbook(); ws=wb.active
 if source=='orders':
  ws.title='Pedidos'; ws.append(['ID Pedido','Nombre Pedido','Cliente','Marca','Modelo','Lote','Procesador','RAM','Disco'])
  for o in CustomerOrder.objects.select_related('customer').order_by('id'): ws.append([o.pk,o.name,o.customer.name,o.brand,o.model,o.lot,o.processor,o.ram,o.disk])
 elif source=='units':
  ws.title='Unidades'; ws.append(['SN','ID Pedido','Pedido','Cliente','Marca','Modelo','Lote','Procesador','RAM','Disco'])
  for u in OrderUnit.objects.select_related('order','order__customer').order_by('id'): ws.append([u.serial_number,u.order_id,u.order.name,u.order.customer.name,u.brand,u.model,u.aiken_lot,u.processor,u.ram,u.disk])
 elif source=='components':
  ws.title='Componentes'; ws.append(['ID','Tipo','Proveedor','Referencia','Fecha','Estado','Observaciones'])
  for c in Component.objects.select_related('supplier').order_by('id'): ws.append([c.pk,c.component_type,c.supplier.name if c.supplier_id else 'Bodega/Stock',c.reference,c.date,c.get_status_display(),c.observations])
 elif source=='rma':
  ws.title='RMA'; ws.append(['ID','Componente','Proveedor','Fecha','Estado','Motivo','Observaciones'])
  for r in RMA.objects.select_related('component','supplier').order_by('id'): ws.append([r.pk,str(r.component),r.supplier.name,r.created_at.replace(tzinfo=None),r.get_status_display(),r.reason,r.observations])
 else: return _deny()
 response=HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'); response['Content-Disposition']=f'attachment; filename="pulsia_{source}.xlsx"'; wb.save(response); return response
